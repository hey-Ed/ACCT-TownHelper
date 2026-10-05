"""Builds the acre catalogue from the category sheets of the CF TownHelper
document (acre listing by frobuddyharry).

Usage: py build_townhelper_data.py "TownHelper.xlsx"
(only needed when the TownHelper document is updated)
"""
import openpyxl, io, re, json, os, sys, hashlib, collections
import numpy as np
from PIL import Image

LIGHT, WATER, SAND, GREY, BRIDGE, LINE, ICON, DARK, UNK = 0, 1, 2, 3, 4, 5, 6, 7, 8
PAL = {(90,170,33): LIGHT, (74,150,33): DARK, (115,199,49): LIGHT, (99,177,41): LIGHT,
       (93,172,35): LIGHT, (8,93,206): WATER, (181,182,181): GREY,
       (255,142,82): BRIDGE, (140,69,8): LINE, (214,215,115): SAND}
CELL = {'FF34A853': DARK, ('t',7): DARK, 'FF5DAC23': LIGHT, 'FF73C731': LIGHT, 'FF4CB61D': LIGHT,
        'FF4285F4': WATER, ('t',4): WATER, 'FF47ACF1': WATER, 'FFFFD966': SAND, 'FFFFE599': SAND,
        'FFE69138': BRIDGE, 'FF999999': GREY}
OBJ = {'FFAA6540': 'tree', 'FFCCCCCC': 'rock', 'FFF3F3F3': 'rock', 'FFB7B7B7': 'rock', 'FFEFEFEF': 'rock'}

def cellkey(c):
    f = c.fill
    if f is None or f.fill_type is None: return None
    fc = f.fgColor
    return fc.rgb if fc.type == 'rgb' else ('t', fc.theme)

def classify_mini(img):
    a = np.array(img.convert('RGB').resize((48, 48), Image.NEAREST)).astype(int)
    keys = np.array(list(PAL.keys())); vals = np.array(list(PAL.values()))
    d = np.abs(a[..., None, :] - keys[None, None]).sum(-1)
    return vals[d.argmin(-1)]

def block_info(ws, R, C):
    tiles = np.full((16, 16), UNK); objs = np.zeros((16, 16), bool)
    for i in range(16):
        for j in range(16):
            k = cellkey(ws.cell(R + i, C + j))
            if k in CELL: tiles[i, j] = CELL[k]
            else: objs[i, j] = True
            if k in OBJ: objs[i, j] = True
    # objects (trees, flowers...): use the surrounding terrain
    for _ in range(6):
        u = tiles == UNK
        if not u.any(): break
        for i, j in zip(*np.nonzero(u)):
            nb = [tiles[y, x] for y, x in ((i-1,j),(i+1,j),(i,j-1),(i,j+1))
                  if 0 <= y < 16 and 0 <= x < 16 and tiles[y, x] != UNK]
            if nb: tiles[i, j] = max(set(nb), key=nb.count)
    tiles[tiles == UNK] = LIGHT
    has_bridge = bool((tiles == BRIDGE).any())
    return tiles, objs, has_bridge

def imghash(p):
    return hashlib.md5(np.array(p.convert('RGB')).tobytes()).hexdigest()

def build(xlsx, out_dir):
    wb = openpyxl.load_workbook(xlsx)
    aa = wb.worksheets[-1]
    # game IDs (All Acres tab), used for display
    idrows = []
    for r in range(13, aa.max_row + 1):
        v = aa.cell(r, 1).value
        if v is None: continue
        s = str(v); m = re.match(r'^(\d+)', s)
        if m and (isinstance(v, (int, float)) or '(' in s):
            idrows.append((r, int(m.group(1)), str(aa.cell(r + 1, 1).value).strip()))
    def raw(ws, R, C): return tuple(cellkey(ws.cell(R+i, C+j)) for i in range(16) for j in range(16))
    aab = {i: (raw(aa, r, 3), n) for r, i, n in idrows}
    by_hash = collections.defaultdict(list)
    for im in aa._images:
        ar = im.anchor._from.row + 1
        rid = [x for x in idrows if x[0] <= ar + 1][-1][1]
        by_hash[imghash(Image.open(io.BytesIO(im._data())))].append(rid)

    os.makedirs(out_dir, exist_ok=True)
    entries = []; arrays = {}
    for sh in ['Buildings', 'CliffsRampsGrasslands', 'Beach', 'Rivers', 'River StartLakes']:
        ws = wb[sh]
        labels = [(r, str(ws.cell(r, 1).value).strip()) for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value]
        imgs = sorted(ws._images, key=lambda im: im.anchor._from.row)
        for li, (lr, lab) in enumerate(labels):
            end = labels[li + 1][0] if li + 1 < len(labels) else ws.max_row + 1
            blocks = sorted((R, c) for R in range(lr, end) for c in range(3, ws.max_column - 14)
                            if cellkey(ws.cell(R, c)) and cellkey(ws.cell(R + 15, c + 15))
                            and (R == 1 or not cellkey(ws.cell(R - 1, c))) and not cellkey(ws.cell(R, c - 1)))
            minis, hashes = [], []
            for im in imgs:
                if lr - 1 <= im.anchor._from.row < end - 1:
                    p = Image.open(io.BytesIO(im._data())).convert('RGB')
                    minis.append(classify_mini(p)); hashes.append(imghash(p))
            for k, (R, C) in enumerate(blocks):
                tiles, objs, bridge = block_info(ws, R, C)
                b48 = np.kron(tiles, np.ones((3, 3), int))
                mh = None
                if len(minis) == len(blocks): ref, how, mh = minis[k], 'mini', hashes[k]
                elif len(minis) == 1: ref, how, mh = minis[0], 'mini', hashes[0]
                else:
                    best, ref, how = 0, b48, 'block'
                    for m, hh in zip(minis, hashes):
                        v = m != LINE
                        sc = (m[v] == b48[v]).mean()
                        if sc > best: best, ref, how, mh = sc, m, 'mini', hh
                r = raw(ws, R, C)
                sims = {i: sum(x == y for x, y in zip(r, ab)) for i, (ab, n) in aab.items()}
                key = f"{sh}|{lab}|{k}"
                e = dict(key=key, sheet=sh, label=lab, k=k, row=R, col=C, sims=sims, mh=mh, bridge=bridge,
                         ramp=any(cellkey(ws.cell(R+i, C+j)) == 'FF73C731' for i in range(16) for j in range(16)),
                         how=how)
                arrays[f"mini_{len(entries)}"] = ref.astype(np.uint8)
                arrays[f"tiles_{len(entries)}"] = tiles.astype(np.uint8)
                arrays[f"objs_{len(entries)}"] = objs
                entries.append(e)
    # game ID: first through the identical All Acres preview image,
    # then split blocks that share a preview (bridge / no bridge)
    for e in entries: e['id'] = None
    groups = collections.defaultdict(list)
    for e in entries:
        if e['mh'] in by_hash: groups[e['mh']].append(e)
    for hh, es in groups.items():
        ids = sorted(set(by_hash[hh]))
        if len(es) == 1 or len(ids) == 1:
            for e in es: e['id'] = max(ids, key=lambda i: e['sims'][i])
        else:
            pairs = sorted(((e['sims'][i], n, i) for n, e in enumerate(es) for i in ids), reverse=True)
            used_e, used_i = set(), set()
            for sc, n, i in pairs:
                if n in used_e or i in used_i: continue
                es[n]['id'] = i; used_e.add(n); used_i.add(i)
    # twins that got the same ID: the less similar one takes the closest free ID
    taken = collections.Counter(e['id'] for e in entries if e['id'] is not None)
    free = set(aab) - set(taken)
    for i, c in list(taken.items()):
        if c < 2: continue
        dup = sorted([e for e in entries if e['id'] == i], key=lambda e: -e['sims'][i])
        for e in dup[1:]:
            if not free: break
            sc, j = max((e['sims'][j], j) for j in free)
            if sc >= 200: e['id'] = j; free.discard(j)
    for e in entries:
        if e['id'] is None and e['mh'] not in by_hash:
            sc, i = max((v, i) for i, v in e['sims'].items())
            if sc >= 238: e['id'] = i
    for e in entries: del e['sims']; del e['mh']
    json.dump(entries, open(os.path.join(out_dir, 'catalog.json'), 'w'), indent=1)
    np.savez_compressed(os.path.join(out_dir, 'catalog.npz'), **arrays)
    return entries

if __name__ == '__main__':
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), 'townhelper_data')
    es = build(sys.argv[1], out)
    print(f"{len(es)} acres saved to {out}")
    sans_id = [e['key'] for e in es if e['id'] is None]
    if sans_id:
        print("Acres without a game ID (missing from All Acres):", ", ".join(sans_id))
